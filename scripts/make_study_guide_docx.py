"""Convert EXAM_STUDY_GUIDE.md to a formatted Word document."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD_PATH = Path(__file__).parent.parent / "EXAM_STUDY_GUIDE.md"
OUT_PATH = Path(__file__).parent.parent / "SoundSight_Study_Guide.docx"


def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)


def add_table_from_md(doc, lines):
    rows = [line for line in lines if line.startswith("|") and "---" not in line]
    if not rows:
        return
    cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
    ncols = len(cells[0])
    table = doc.add_table(rows=len(cells), cols=ncols)
    table.style = "Table Grid"
    for i, row_data in enumerate(cells):
        for j, cell_text in enumerate(row_data):
            cell = table.cell(i, j)
            cell.text = cell_text
            run = cell.paragraphs[0].runs
            if i == 0 and run:
                run[0].bold = True
                run[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if i == 0:
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "2E4057")
                tcPr.append(shd)
    doc.add_paragraph()


def apply_inline(run_text, para):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", run_text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            r = para.add_run(part[1:-1])
            r.font.name = "Courier New"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xC7, 0x25, 0x4F)
        elif part.startswith("**") and part.endswith("**"):
            r = para.add_run(part[2:-2])
            r.bold = True
        else:
            para.add_run(part)


def convert(md_path: Path, out_path: Path):
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # Title page
    title = doc.add_heading("SoundSight", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Exam Study Guide")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(14)
    sub.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_paragraph()
    doc.add_page_break()

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines = []
    table_lines = []

    while i < len(lines):
        line = lines[i]

        # Code block
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                p = doc.add_paragraph()
                p.style = doc.styles["Normal"]
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run("\n".join(code_lines))
                r.font.name = "Courier New"
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "F0F0F0")
                p._p.get_or_add_pPr().append(shd)
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Table
        if line.startswith("|"):
            table_lines.append(line)
            i += 1
            continue
        else:
            if table_lines:
                add_table_from_md(doc, table_lines)
                table_lines = []

        # Headings
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("---"):
            add_horizontal_rule(doc)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            apply_inline(line[2:], p)
        elif line.strip() == "":
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            apply_inline(line, p)

        i += 1

    if table_lines:
        add_table_from_md(doc, table_lines)

    doc.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    convert(MD_PATH, OUT_PATH)
